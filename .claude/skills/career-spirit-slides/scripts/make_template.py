#!/usr/bin/env python3
"""
make_template.py
SOURCE.pptx (cw_DS_source.pptx, 39枚) からテンプレートスライド＋プリセットスライドを
抽出して TEMPLATE.pptx を生成するスクリプト。

【アルゴリズム】jcc-slides 同様
  1. SOURCE.pptx を1インスタンスだけ開く
  2. 対象スライドを同一 prs 内で末尾に複製（メディア重複を防ぐ）
  3. 元の39枚を後ろから削除

使い方:
  python make_template.py \
    --input  /path/to/SOURCE.pptx \
    --output /path/to/TEMPLATE.pptx

生成後のインデックス対応（24枚）:
  ── ベーステンプレート（9枚）──
  [0] title           [1] section_dark    [2] section_check
  [3] title_light     [4] question        [5] answer
  [6] content_list_4  [7] memo            [8] profile

  ── fixed プリセット（10枚）──
  [9]  be_do_have       [10] iceberg          [11] loop_diagram
  [12] trap_diagram     [13] walls4           [14] goal_brake
  [15] step3_cards      [16] timeline_kpi     [17] group_phases
  [18] case_study

  ── fillable プリセット（5枚）──
  [19] flow4_horizontal [20] return_card      [21] divided_world
  [22] stats_card       [23] before_after
"""

import sys
import copy
import argparse
import os

try:
    from pptx import Presentation
    from lxml import etree
except ImportError:
    print("python-pptx / lxml が必要です: pip install python-pptx lxml --break-system-packages")
    sys.exit(1)

# 抽出対象: (元インデックス, slide_type)
EXTRACT_SLIDES = [
    # ── ベーステンプレート（9枚） ──
    (0,  "title"),
    (1,  "section_dark"),
    (2,  "section_check"),
    (3,  "title_light"),
    (4,  "question"),
    (5,  "answer"),
    (6,  "content_list_4"),
    (10, "memo"),
    (37, "profile"),
    # ── fixed プリセット（10枚） ──
    (8,  "be_do_have"),
    (22, "iceberg"),
    (26, "loop_diagram"),
    (28, "trap_diagram"),
    (30, "walls4"),
    (31, "goal_brake"),
    (33, "step3_cards"),
    (34, "timeline_kpi"),
    (36, "group_phases"),
    (38, "case_study"),
    # ── fillable プリセット（5枚） ──
    (7,  "flow4_horizontal"),
    (16, "return_card"),
    (27, "divided_world"),
    (24, "stats_card"),
    (35, "before_after"),
]

NS_R    = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
R_EMBED = f"{{{NS_R}}}embed"
R_ID    = f"{{{NS_R}}}id"
R_HREF  = f"{{{NS_R}}}href"


def duplicate_within(prs, src_idx):
    """prs の src_idx 番スライドを同一 prs 内で複製して末尾に追加する。"""
    template = prs.slides[src_idx]
    new_slide = prs.slides.add_slide(template.slide_layout)

    rId_map = {}
    for rel in template.part.rels.values():
        reltype_short = rel.reltype.split("/")[-1]
        if reltype_short in ("slideLayout", "notesSlide"):
            continue
        if rel.is_external:
            new_rId = new_slide.part.target_ref(rel.reltype, rel.target_ref)
        else:
            new_rId = new_slide.part.relate_to(rel.target_part, rel.reltype)
        rId_map[rel.rId] = new_rId

    tree_copy = copy.deepcopy(template.shapes._spTree)
    for elem in tree_copy.iter():
        for attr in (R_EMBED, R_ID, R_HREF):
            if attr in elem.attrib and elem.attrib[attr] in rId_map:
                elem.attrib[attr] = rId_map[elem.attrib[attr]]

    dst_tree = new_slide.shapes._spTree
    for child in list(dst_tree):
        dst_tree.remove(child)
    for child in list(tree_copy):
        dst_tree.append(child)

    return new_slide


def delete_slide(prs, idx):
    """Remove slide entry from sldIdLst AND drop the relationship.
    Without dropping the rel, the part may stay in the package as an orphan
    causing duplicate-name issues on subsequent generation."""
    sldIdLst = prs.slides._sldIdLst
    items = list(sldIdLst)
    target = items[idx]
    rId = target.rId
    prs.part.drop_rel(rId)
    sldIdLst.remove(target)


def main():
    parser = argparse.ArgumentParser(description="SOURCE.pptx からテンプレートスライドを抽出")
    parser.add_argument("--input",  default=None, help="入力 PPTX パス（デフォルト: assets/SOURCE.pptx）")
    parser.add_argument("--output", default=None, help="出力 PPTX パス（デフォルト: assets/TEMPLATE.pptx）")
    args = parser.parse_args()

    script_dir = os.path.dirname(os.path.abspath(__file__))
    assets_dir = os.path.join(script_dir, "..", "assets")

    input_path  = args.input  or os.path.join(assets_dir, "SOURCE.pptx")
    output_path = args.output or os.path.join(assets_dir, "TEMPLATE.pptx")

    if not os.path.exists(input_path):
        print(f"エラー: 入力ファイルが見つかりません: {input_path}")
        sys.exit(1)

    print(f"入力: {input_path}")
    print(f"出力: {output_path}")

    prs = Presentation(input_path)
    original_count = len(prs.slides)
    print(f"元スライド数: {original_count}")

    extract_indices = [idx for idx, _ in EXTRACT_SLIDES]
    print(f"抽出するスライド（0始まり）: {extract_indices}\n")

    # STEP 1: 対象スライドを末尾に複製
    for orig_idx, slide_type in EXTRACT_SLIDES:
        print(f"  複製: slide[{orig_idx:2d}] → {slide_type}")
        duplicate_within(prs, orig_idx)

    # STEP 2: 元のスライドを削除（後ろから）
    print(f"\n元の {original_count} 枚を削除中...")
    for i in range(original_count - 1, -1, -1):
        delete_slide(prs, i)

    prs.save(output_path)

    src_size  = os.path.getsize(input_path)  / 1024 / 1024
    dst_size  = os.path.getsize(output_path) / 1024 / 1024
    final_count = len(Presentation(output_path).slides)

    print(f"\n完了: {output_path}")
    print(f"最終スライド数: {final_count} 枚")
    print(f"ファイルサイズ: {src_size:.1f} MB → {dst_size:.1f} MB")

    print("\n── 新インデックス対応表 ──")
    for new_idx, (orig_idx, slide_type) in enumerate(EXTRACT_SLIDES):
        print(f"  [{new_idx:2d}] {slide_type:20s}  ← 元[{orig_idx}]")


if __name__ == "__main__":
    main()

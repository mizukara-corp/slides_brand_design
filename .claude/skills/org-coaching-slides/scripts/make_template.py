#!/usr/bin/env python3
"""
make_template.py
training3.pptx（新ブランド・ベースデザイン）と training2.pptx（図解スライド）から
組織コーチング用 TEMPLATE.pptx を抽出生成するスクリプト。

【アルゴリズム】
  1. training3.pptx → 末尾に必要な base スライドを複製
  2. training3.pptx に training2.pptx の図解スライドを差し込む
     （同一 prs 内でメディア重複させない複製ロジックを 2 ファイル間で拡張）
  3. 元の training3 スライドを後ろから全削除
  4. 結果を assets/TEMPLATE.pptx として保存

生成後のインデックス対応（20枚）:
  ── ベーステンプレート（training3 由来 / 6枚） ──
  [ 0] title           ← training3[ 0]  Organizational Coaching / DAY◯
  [ 1] meditation      ← training3[ 1]  瞑想（センター・ダーク）
  [ 2] leader_message  ← training3[ 2]  Leader's Message リーダーメッセージ
  [ 3] check_in        ← training3[ 3]  Check in + 問いかけ
  [ 4] rest_time       ← training3[ 9]  Rest Time（カラフルグラデ）
  [ 5] closing_light   ← training3[11]  ライト背景クロージング

  ── テキスト編集可（training3 由来・プレースホルダ付き / 2枚） ──
  [ 6] question        ← training3[ 4]  Q. + 問い本文
  [ 7] purpose         ← training3[ 6]  Purpose + 段落本文

  ── テキスト編集可（training2 由来 / 7枚） ──
  [ 8] key_point       ← training2[12]  重要な観点フレーム＋メッセージ
  [ 9] section_divider ← training2[19]  章タイトル（中央配置）
  [10] notice          ← training2[20]  DAY告知（フレーム＋番号付きリスト）
  [11] agenda_h0       ← training2[ 0]  アジェンダ（強調なし）
  [12] agenda_h1       ← training2[ 1]  アジェンダ（item1 強調）
  [13] agenda_h2       ← training2[ 2]  アジェンダ（item2 強調）
  [14] agenda_h3       ← training2[ 3]  アジェンダ（item3 強調）

  ── fillable プリセット（training2 由来 / 5枚） ──
  [15] result_chain    ← training2[ 9]  4ボックス連鎖（結果←行動←感情←信念）
  [16] compare_2card   ← training2[11]  2カード比較（DO/BE）
  [17] compare_3card   ← training2[17]  3カード比較（一致/合意/了解）
  [18] loop_learning   ← training2[15]  3円ループ図（前提→行動→結果）
  [19] circle_overlap  ← training2[18]  中央大円＋3サテライト円
"""

import sys
import copy
import argparse
import os

try:
    from pptx import Presentation
    from lxml import etree
except ImportError:
    print("python-pptx / lxml が必要です: pip install python-pptx lxml --break-system-packages")
    sys.exit(1)


# (source_key, source_idx, slide_type) — 上から順に末尾に追加される
EXTRACT_PLAN = [
    # ── training3 由来（base デザイン）──
    ("training3", 0,  "title"),
    ("training3", 1,  "meditation"),
    ("training3", 2,  "leader_message"),
    ("training3", 3,  "check_in"),
    ("training3", 9,  "rest_time"),
    ("training3", 11, "closing_light"),
    ("training3", 4,  "question"),
    ("training3", 6,  "purpose"),
    # ── training2 由来（図解・コンテンツ）──
    ("training2", 12, "key_point"),
    ("training2", 19, "section_divider"),
    ("training2", 20, "notice"),
    ("training2", 0,  "agenda_h0"),
    ("training2", 1,  "agenda_h1"),
    ("training2", 2,  "agenda_h2"),
    ("training2", 3,  "agenda_h3"),
    ("training2", 9,  "result_chain"),
    ("training2", 11, "compare_2card"),
    ("training2", 17, "compare_3card"),
    ("training2", 15, "loop_learning"),
    ("training2", 18, "circle_overlap"),
]

NS_R    = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
R_EMBED = f"{{{NS_R}}}embed"
R_ID    = f"{{{NS_R}}}id"
R_HREF  = f"{{{NS_R}}}href"


def duplicate_within(prs, src_idx):
    """prs の src_idx 番スライドを同一 prs 内で複製して末尾に追加する。"""
    template = prs.slides[src_idx]
    new_slide = prs.slides.add_slide(template.slide_layout)

    rId_map = {}
    for rel in template.part.rels.values():
        reltype_short = rel.reltype.split("/")[-1]
        if reltype_short in ("slideLayout", "notesSlide"):
            continue
        if rel.is_external:
            new_rId = new_slide.part.target_ref(rel.reltype, rel.target_ref)
        else:
            new_rId = new_slide.part.relate_to(rel.target_part, rel.reltype)
        rId_map[rel.rId] = new_rId

    tree_copy = copy.deepcopy(template.shapes._spTree)
    for elem in tree_copy.iter():
        for attr in (R_EMBED, R_ID, R_HREF):
            if attr in elem.attrib and elem.attrib[attr] in rId_map:
                elem.attrib[attr] = rId_map[elem.attrib[attr]]

    dst_tree = new_slide.shapes._spTree
    for child in list(dst_tree):
        dst_tree.remove(child)
    for child in list(tree_copy):
        dst_tree.append(child)

    return new_slide


def import_slide_from(dst_prs, src_prs, src_idx):
    """別 prs (src_prs) の src_idx 番スライドを dst_prs に追加する。
    src スライドの slideLayout 名と同じレイアウトを dst_prs から探して紐付ける
    （training3 と training2 はレイアウト名を共有しているため、これで装飾が一致する）。
    一致しない場合は dst_prs のレイアウト 0 にフォールバック。
    """
    src_slide = src_prs.slides[src_idx]
    src_layout_name = src_slide.slide_layout.name

    # dst_prs から同名レイアウトを検索
    matched_layout = None
    for lay in dst_prs.slide_layouts:
        if lay.name == src_layout_name:
            matched_layout = lay
            break
    if matched_layout is None:
        matched_layout = dst_prs.slide_layouts[0]
        print(f"  ⚠ layout '{src_layout_name}' が dst にないため layout[0] を使用")

    new_slide = dst_prs.slides.add_slide(matched_layout)

    # --- 関連リソース (image / chart / etc.) を dst 側に移植 ---
    rId_map = {}
    for rel in src_slide.part.rels.values():
        reltype_short = rel.reltype.split("/")[-1]
        if reltype_short in ("slideLayout", "notesSlide"):
            continue
        if rel.is_external:
            new_rId = new_slide.part.target_ref(rel.reltype, rel.target_ref)
        else:
            new_rId = new_slide.part.relate_to(rel.target_part, rel.reltype)
        rId_map[rel.rId] = new_rId

    # --- 図形ツリーを丸ごとコピーして rId を新 ID に張替え ---
    tree_copy = copy.deepcopy(src_slide.shapes._spTree)
    for elem in tree_copy.iter():
        for attr in (R_EMBED, R_ID, R_HREF):
            if attr in elem.attrib and elem.attrib[attr] in rId_map:
                elem.attrib[attr] = rId_map[elem.attrib[attr]]

    dst_tree = new_slide.shapes._spTree
    for child in list(dst_tree):
        dst_tree.remove(child)
    for child in list(tree_copy):
        dst_tree.append(child)

    # --- スライド背景 (cSld/bg) を src から dst に複製 ---
    src_cSld = src_slide._element.find(
        "{http://schemas.openxmlformats.org/presentationml/2006/main}cSld"
    )
    dst_cSld = new_slide._element.find(
        "{http://schemas.openxmlformats.org/presentationml/2006/main}cSld"
    )
    if src_cSld is not None and dst_cSld is not None:
        src_bg = src_cSld.find(
            "{http://schemas.openxmlformats.org/presentationml/2006/main}bg"
        )
        # 既存 bg を除去
        for old_bg in dst_cSld.findall(
            "{http://schemas.openxmlformats.org/presentationml/2006/main}bg"
        ):
            dst_cSld.remove(old_bg)
        if src_bg is not None:
            new_bg = copy.deepcopy(src_bg)
            for elem in new_bg.iter():
                for attr in (R_EMBED, R_ID, R_HREF):
                    if attr in elem.attrib and elem.attrib[attr] in rId_map:
                        elem.attrib[attr] = rId_map[elem.attrib[attr]]
            # cSld 内では bg は spTree より前に置く
            dst_cSld.insert(0, new_bg)

    return new_slide


def delete_slide(prs, idx):
    sldIdLst = prs.slides._sldIdLst
    items = list(sldIdLst)
    target = items[idx]
    rId = target.rId
    prs.part.drop_rel(rId)
    sldIdLst.remove(target)


def main():
    parser = argparse.ArgumentParser(description="training3 + training2 から TEMPLATE.pptx を生成")
    parser.add_argument("--training3", default=None, help="training3.pptx へのパス")
    parser.add_argument("--training2", default=None, help="training2.pptx へのパス")
    parser.add_argument("--output", default=None, help="出力 TEMPLATE.pptx へのパス")
    args = parser.parse_args()

    script_dir = os.path.dirname(os.path.abspath(__file__))
    skill_dir = os.path.abspath(os.path.join(script_dir, ".."))
    assets_dir = os.path.join(skill_dir, "assets")

    src3_path = args.training3 or os.path.join(skill_dir, "training3.pptx")
    src2_path = args.training2 or os.path.join(skill_dir, "training2.pptx")
    output_path = args.output or os.path.join(assets_dir, "TEMPLATE.pptx")

    for p in (src3_path, src2_path):
        if not os.path.exists(p):
            print(f"エラー: 入力ファイルが見つかりません: {p}")
            sys.exit(1)

    print(f"training3: {src3_path}")
    print(f"training2: {src2_path}")
    print(f"出力       : {output_path}")

    # training3 をベースに開き、必要なスライドを末尾に複製していく
    dst_prs = Presentation(src3_path)
    src2_prs = Presentation(src2_path)
    original_count = len(dst_prs.slides)
    print(f"\ntraining3 スライド数: {original_count}")
    print(f"training2 スライド数: {len(src2_prs.slides)}\n")

    for src_key, src_idx, slide_type in EXTRACT_PLAN:
        if src_key == "training3":
            print(f"  複製: training3[{src_idx:2d}] → {slide_type}")
            duplicate_within(dst_prs, src_idx)
        elif src_key == "training2":
            print(f"  輸入: training2[{src_idx:2d}] → {slide_type}")
            import_slide_from(dst_prs, src2_prs, src_idx)
        else:
            raise ValueError(f"unknown source key: {src_key}")

    # 元 training3 のスライドを後ろから全削除
    print(f"\n元 training3 の {original_count} 枚を削除中...")
    for i in range(original_count - 1, -1, -1):
        delete_slide(dst_prs, i)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    dst_prs.save(output_path)

    src3_size = os.path.getsize(src3_path) / 1024 / 1024
    dst_size = os.path.getsize(output_path) / 1024 / 1024
    final_count = len(Presentation(output_path).slides)

    print(f"\n完了: {output_path}")
    print(f"最終スライド数: {final_count} 枚")
    print(f"ファイルサイズ: {src3_size:.1f} MB → {dst_size:.1f} MB")

    print("\n── 新インデックス対応表 ──")
    for new_idx, (src_key, src_idx, slide_type) in enumerate(EXTRACT_PLAN):
        print(f"  [{new_idx:2d}] {slide_type:18s}  ← {src_key}[{src_idx}]")


if __name__ == "__main__":
    main()
